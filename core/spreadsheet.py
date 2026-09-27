import csv
import json
import math
import os
import tempfile
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook


class SpreadsheetError(ValueError):
    """Raised when a tabular import or export cannot be completed."""


class SpreadsheetStore:
    VALID_COLUMN_TYPES = {"text", "number", "date"}
    MAX_IMPORT_BYTES = 50 * 1024 * 1024
    MAX_IMPORT_ROWS = 10_000
    MAX_IMPORT_COLUMNS = 64

    def __init__(self, draft_path=None):
        if draft_path is None:
            root = Path.home() / "Library" / "Application Support" / "GoScan"
            draft_path = root / "table_draft.json"
        self.draft_path = Path(draft_path)

    @staticmethod
    def validate_value(value, column_type):
        if column_type not in SpreadsheetStore.VALID_COLUMN_TYPES:
            raise SpreadsheetError(f"Tipe kolom tidak dikenal: {column_type}")
        if value is None or str(value).strip() == "":
            return True
        if column_type == "text":
            return True
        if column_type == "number":
            try:
                return math.isfinite(float(str(value).replace(",", ".")))
            except ValueError:
                return False
        try:
            if isinstance(value, (date, datetime)):
                return True
            text = str(value).strip()
            try:
                datetime.strptime(text, "%Y-%m-%d")
                return True
            except ValueError:
                datetime.fromisoformat(text)
                return True
        except ValueError:
            return False

    @staticmethod
    def import_file(file_path):
        path = Path(file_path)
        suffix = path.suffix.lower()
        if suffix not in {".csv", ".xlsx", ".xlsm"}:
            raise SpreadsheetError("Format impor hanya mendukung .csv, .xlsx, atau .xlsm.")
        if path.stat().st_size > SpreadsheetStore.MAX_IMPORT_BYTES:
            raise SpreadsheetError("Ukuran file impor melebihi batas 50 MB.")
        if suffix == ".csv":
            with path.open("r", newline="", encoding="utf-8-sig") as source:
                sample = source.read(4096)
                source.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
                except csv.Error:
                    dialect = csv.excel
                rows = []
                for row in csv.reader(source, dialect):
                    if len(row) > SpreadsheetStore.MAX_IMPORT_COLUMNS:
                        raise SpreadsheetError("File memiliki lebih dari 64 kolom.")
                    rows.append(row)
                    if len(rows) > SpreadsheetStore.MAX_IMPORT_ROWS + 1:
                        raise SpreadsheetError("File memiliki lebih dari 10.000 baris data.")
        elif suffix in {".xlsx", ".xlsm"}:
            try:
                workbook = load_workbook(path, read_only=True, data_only=True)
            except Exception as error:
                raise SpreadsheetError(f"File Excel tidak valid atau rusak: {error}") from error
            try:
                sheet = workbook.active
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    if len(row) > SpreadsheetStore.MAX_IMPORT_COLUMNS:
                        raise SpreadsheetError("File memiliki lebih dari 64 kolom.")
                    rows.append(list(row))
                    if len(rows) > SpreadsheetStore.MAX_IMPORT_ROWS + 1:
                        raise SpreadsheetError("File memiliki lebih dari 10.000 baris data.")
            except SpreadsheetError:
                raise
            except Exception as error:
                raise SpreadsheetError(f"Data Excel tidak dapat dibaca: {error}") from error
            finally:
                workbook.close()
        else:
            raise SpreadsheetError("Format impor hanya mendukung .csv, .xlsx, atau .xlsm.")

        while rows and not any(value not in (None, "") for value in rows[-1]):
            rows.pop()
        if not rows:
            return [], []
        headers = [str(value) if value is not None else "" for value in rows[0]]
        if not any(header.strip() for header in headers):
            raise SpreadsheetError("Baris pertama file harus memiliki judul kolom.")
        headers = [header.strip() or f"Kolom {index + 1}" for index, header in enumerate(headers)]
        data = []
        for row in rows[1:]:
            padded = row[:len(headers)] + [""] * max(0, len(headers) - len(row))
            normalized = []
            for value in padded:
                if value is None:
                    value = ""
                elif isinstance(value, datetime) and value.time().isoformat() == "00:00:00":
                    value = value.date().isoformat()
                elif isinstance(value, (date, datetime)):
                    value = value.isoformat()
                normalized.append(value)
            data.append(normalized)
        return headers, data

    @staticmethod
    def export_file(file_path, headers, rows):
        path = Path(file_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        suffix = path.suffix.lower()
        if suffix not in {".csv", ".xlsx"}:
            raise SpreadsheetError("Format ekspor hanya mendukung .csv atau .xlsx.")
        rows = [list(row) for row in rows]
        while rows and not any(value not in (None, "") for value in rows[-1]):
            rows.pop()

        file_descriptor, temporary_path = tempfile.mkstemp(
            prefix=f"{path.stem}_", suffix=suffix, dir=path.parent
        )
        os.close(file_descriptor)
        try:
            if suffix == ".csv":
                with open(temporary_path, "w", newline="", encoding="utf-8-sig") as destination:
                    writer = csv.writer(destination)
                    writer.writerow([SpreadsheetStore._csv_safe_value(value) for value in headers])
                    writer.writerows(
                        [SpreadsheetStore._csv_safe_value(value) for value in row]
                        for row in rows
                    )
                    destination.flush()
                    os.fsync(destination.fileno())
            else:
                workbook = Workbook()
                sheet = workbook.active
                sheet.title = "Data"
                sheet.append(list(headers))
                for column_index, value in enumerate(headers, start=1):
                    if isinstance(value, str) and value.startswith("="):
                        sheet.cell(1, column_index).data_type = "s"
                for row in rows:
                    sheet.append(list(row))
                    for column_index, value in enumerate(row, start=1):
                        if isinstance(value, str) and value.startswith("="):
                            sheet.cell(sheet.max_row, column_index).data_type = "s"
                for column_cells in sheet.columns:
                    width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 60)
                    sheet.column_dimensions[column_cells[0].column_letter].width = max(12, width)
                workbook.save(temporary_path)
                workbook.close()
            os.replace(temporary_path, path)
        except Exception:
            try:
                os.unlink(temporary_path)
            except OSError:
                pass
            raise

    @staticmethod
    def _csv_safe_value(value):
        if not isinstance(value, str):
            return value
        stripped = value.lstrip()
        if stripped.startswith(("=", "+", "@")):
            return "'" + value
        if stripped.startswith("-"):
            try:
                float(stripped.replace(",", "."))
            except ValueError:
                return "'" + value
        return value

    def save_draft(self, headers, rows, column_types):
        payload = {
            "version": 1,
            "headers": list(headers),
            "rows": [list(row) for row in rows],
            "column_types": list(column_types),
            "saved_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        }
        self.draft_path.parent.mkdir(parents=True, exist_ok=True)
        file_descriptor, temporary_path = tempfile.mkstemp(
            prefix=f"{self.draft_path.name}.", suffix=".tmp", dir=self.draft_path.parent
        )
        try:
            with os.fdopen(file_descriptor, "w", encoding="utf-8") as destination:
                json.dump(payload, destination, ensure_ascii=False, indent=2)
                destination.flush()
                os.fsync(destination.fileno())
            os.replace(temporary_path, self.draft_path)
        except Exception:
            try:
                os.unlink(temporary_path)
            except OSError:
                pass
            raise

    def load_draft(self):
        if not self.draft_path.is_file():
            return None
        try:
            with self.draft_path.open("r", encoding="utf-8") as source:
                payload = json.load(source)
            if payload.get("version") != 1:
                raise SpreadsheetError("Versi draf tidak didukung.")
            headers = payload["headers"]
            rows = payload["rows"]
            types = payload["column_types"]
            if not isinstance(headers, list) or not isinstance(rows, list):
                raise SpreadsheetError("Struktur draf tidak valid.")
            if not headers or len(headers) > self.MAX_IMPORT_COLUMNS:
                raise SpreadsheetError("Jumlah kolom draf tidak valid.")
            if len(rows) > self.MAX_IMPORT_ROWS:
                raise SpreadsheetError("Jumlah baris draf melebihi batas aplikasi.")
            if len(types) != len(headers) or any(t not in self.VALID_COLUMN_TYPES for t in types):
                raise SpreadsheetError("Jenis kolom pada draf tidak valid.")
            normalized_rows = []
            for row in rows:
                values = list(row[:len(headers)])
                values.extend([""] * (len(headers) - len(values)))
                normalized_rows.append(values)
            return {"headers": headers, "rows": normalized_rows, "column_types": types}
        except (OSError, json.JSONDecodeError, KeyError, TypeError) as error:
            raise SpreadsheetError(f"Draf lokal tidak dapat dipulihkan: {error}") from error

    def clear_draft(self):
        try:
            self.draft_path.unlink(missing_ok=True)
        except OSError as error:
            raise SpreadsheetError(f"Draf lokal tidak dapat dihapus: {error}") from error
